import io, re
from pathlib import Path
import cv2, joblib, numpy as np, pandas as pd, streamlit as st
from PIL import Image
from skimage.feature import hog
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

B=Path(__file__).parent
st.set_page_config(page_title='分類モデルOCR→Excel',page_icon='🔢',layout='wide')
st.title('🔢 一桁分類モデルで測定値を認識')
st.caption('正解付き画像から学習した0〜9分類モデルを使います。小数点は最後の桁の前へ自動挿入します。')
MODEL=joblib.load(B/'digit_model.joblib')
SLOTS=[(0,39),(35,74),(82,124)]
DECIMAL_GAP=(74,82)

def timekey(name):
 m=re.search(r'_(\d+(?:\.\d+)?)s(?:\.[^.]+)?$',name,re.I)
 return (0,float(m.group(1))) if m else (1,name.lower())

def crop_aligned(rgb):
 h,w=rgb.shape[:2]
 cx=int(round(w*.505));cy=int(round(h*.544));size=max(20,int(round(min(h,w)*.072)));u=size/30
 box=(max(0,int(cx-48*u)),max(0,int(cy-33*u)),min(w,int(cx-1*u)),min(h,int(cy-3*u)))
 region=rgb[box[1]:box[3],box[0]:box[2]]
 if region.size==0:return None,box,(cx,cy)
 g=cv2.cvtColor(region,cv2.COLOR_RGB2GRAY)
 return cv2.resize(g,(126,72),interpolation=cv2.INTER_CUBIC),box,(cx,cy)

def glyph(aligned,pos):
 x1,x2=SLOTS[pos]
 q=aligned[2:70,x1:x2]
 q=cv2.resize(q,(32,48),interpolation=cv2.INTER_CUBIC)
 return cv2.createCLAHE(2.0,(4,4)).apply(q)

def feature(q):
 return hog(q,orientations=9,pixels_per_cell=(8,8),cells_per_block=(2,2),block_norm='L2-Hys')

def recognize(rgb):
 aligned,box,center=crop_aligned(rgb)
 if aligned is None:return '',0.0,box,center,[],None,[]
 gs=[glyph(aligned,i) for i in range(3)]
 X=np.asarray([feature(g) for g in gs])
 pred=MODEL.predict(X)
 prob=MODEL.predict_proba(X)
 classes=MODEL.classes_
 details=[];conf=[]
 for row in prob:
  order=np.argsort(row)[::-1][:4]
  details.append([(str(classes[j]),float(row[j])) for j in order])
  conf.append(float(row[order[0]]))
 raw=''.join(pred)
 return raw[:2]+'.'+raw[2],float(np.mean(conf)),box,center,details,aligned,gs

def cellok(value):return bool(re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]*',value.strip()))

excel=st.file_uploader('1. 入力先Excel',type=['xlsx'])
files=sorted(st.file_uploader('2. 写真を選択',type=['png','jpg','jpeg','webp'],accept_multiple_files=True) or [],key=lambda f:timekey(f.name))
if files:st.info('時間順：'+' → '.join(f.name for f in files))
rows=[]
for i,f in enumerate(files):
 rgb=np.array(Image.open(f).convert('RGB'))
 value,conf,box,center,details,aligned,gs=recognize(rgb)
 marked=rgb.copy()
 if box:cv2.rectangle(marked,(box[0],box[1]),(box[2],box[3]),(0,255,0),2)
 if center:cv2.drawMarker(marked,center,(255,0,255),cv2.MARKER_CROSS,20,2)
 left,right=st.columns([1,2])
 left.image(marked,caption=f.name,width='stretch')
 value=right.text_input('認識結果',value,key=f'{i}_{f.name}')
 right.write(f'平均分類確率：{conf:.2f}')
 if conf<.55:right.warning('確率が低いため認識結果を確認してください。')
 if aligned is not None:
  preview=cv2.cvtColor(aligned,cv2.COLOR_GRAY2RGB)
  for x1,x2 in SLOTS:cv2.rectangle(preview,(x1,2),(x2,70),(255,0,0),1)
  cv2.rectangle(preview,(DECIMAL_GAP[0],48),(DECIMAL_GAP[1],70),(255,255,0),1)
  with right.expander('認識に使った画像と候補'):
   st.image(preview,caption='赤枠＝数字、黄色枠＝除外する小数点',width='stretch')
   st.image(gs,caption=['1桁目','2桁目','3桁目'],width=90)
   for k,cands in enumerate(details):st.write(f'{k+1}桁目：'+', '.join(f'{d}({p:.2f})' for d,p in cands))
 rows.append({'filename':f.name,'value':value,'confidence':round(conf,3)})
if rows:
 st.dataframe(pd.DataFrame(rows),width='stretch')
 if excel:
  wb=load_workbook(io.BytesIO(excel.getvalue()));sn=st.selectbox('入力シート',wb.sheetnames);start=st.text_input('開始セル','C4')
  if cellok(start):
   letters,r0=coordinate_from_string(start.upper());col=column_index_from_string(letters);ws=wb[sn];invalid=[]
   for j,row in enumerate(rows):
    try:
     cell=ws.cell(r0+j,col,float(row['value']));cell.number_format='0.0'
    except (ValueError,TypeError):invalid.append(row['filename'])
   if invalid:st.error('数値に変換できない画像：'+'、'.join(invalid))
   else:
    end=f'{get_column_letter(col)}{r0+len(rows)-1}';out=io.BytesIO();wb.save(out)
    st.success(f'{start.upper()}:{end}へ入力しました')
    st.download_button('入力済みExcelをダウンロード',out.getvalue(),f'{Path(excel.name).stem}_入力済み.xlsx')
